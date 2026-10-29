"""Write laid-out items into a .pptx with python-pptx.

Everything stays editable: text is real text in the theme font, shapes are
PowerPoint shapes, charts are native charts with their data inside, and an SVG
picture keeps the SVG itself next to its PNG fallback so PowerPoint can turn it
into shapes (Convert to Shape).
"""

from __future__ import annotations

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE, XL_TICK_MARK
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

from fonts import parse_runs
from layout import H, W

ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
SVG_EXT_URI = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"
SVG_NS = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
NO_TABLE_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"


def nice_axis(vmin: float, vmax: float, ticks: int = 4) -> tuple[float, float, float]:
    """Axis bounds on round numbers with about `ticks` gridlines, a little room above and below."""
    import math

    span = (vmax - vmin) or abs(vmax) or 1.0
    raw = span / ticks
    mag = 10 ** math.floor(math.log10(raw))
    step = next(m * mag for m in (1, 2, 2.5, 5, 10) if m * mag >= raw)
    lo = math.floor((vmin - span * 0.15) / step) * step
    if vmin >= 0:
        lo = max(0.0, lo)
    hi = math.ceil((vmax + span * 0.15) / step) * step
    return lo, hi, step


class Writer:
    def __init__(self, theme: dict, ctx, logo_dirs: list[Path], svg_to_png, cache: Path):
        self.theme = theme
        self.ctx = ctx
        self.cache = cache
        self.family = theme.get("fonts", {}).get("family", "Pretendard")
        self.logo_dirs = logo_dirs
        self.svg_to_png = svg_to_png
        self.prs = Presentation()
        self.prs.slide_width = Pt(W)
        self.prs.slide_height = Pt(H)
        self.blank = self.prs.slide_layouts[6]

    # colors ------------------------------------------------------------------
    def hex(self, token: str | None) -> str | None:
        if token is None:
            return None
        colors = self.theme["colors"]
        if token in colors:
            return colors[token]
        self.ctx.err("color", f"테마에 없는 색 이름입니다: {token!r} (쓸 수 있는 이름: {', '.join(colors)})")
        return colors["text"]

    def rgb(self, token):
        return RGBColor.from_string(self.hex(token).lstrip("#"))

    # slide -------------------------------------------------------------------
    def add_slide(self, items: list, notes: str = ""):
        slide = self.prs.slides.add_slide(self.blank)
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = self.rgb("bg")
        for it in items:
            t = it["t"]
            if t == "text":
                self.text(slide, it)
            elif t in ("rect", "round", "oval"):
                self.shape(slide, it)
            elif t == "line":
                self.line(slide, it)
            elif t == "image":
                self.image(slide, it)
            elif t == "chart":
                self.chart(slide, it)
            elif t == "table":
                self.table(slide, it)
            elif t == "logo":
                self.logo(slide, it)
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        return slide

    def save(self, path: Path):
        self.prs.save(str(path))

    # text ----------------------------------------------------------------------
    def set_font(self, rPr):
        rPr.set("lang", "ko-KR")
        rPr.set("altLang", "en-US")
        for tag in ("a:latin", "a:ea", "a:cs"):
            el = rPr.find(qn(tag))
            if el is None:
                el = etree.SubElement(rPr, qn(tag))
            el.set("typeface", self.family)

    def text(self, slide, it):
        x, y, w, h = it["box"]
        tb = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
        tb.name = it["role"]
        tf = tb.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.TOP
        color, accent = it["color"], it.get("accent", "accent")
        for i, para in enumerate(it["paras"] or [""]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = ALIGN.get(it["align"], PP_ALIGN.LEFT)
            p.line_spacing = Pt(it["size"] * it["lh"])
            if i and it.get("para_gap"):
                p.space_before = Pt(it["para_gap"])
            for chunk, bold, acc in parse_runs(para):
                r = p.add_run()
                r.text = chunk
                r.font.size = Pt(it["size"])
                r.font.bold = bool(it["bold"] or bold)
                r.font.color.rgb = self.rgb(accent if acc else color)
                self.set_font(r._r.get_or_add_rPr())
            if it.get("bullets"):
                self.bullet(p, it["indent"], accent)

    def bullet(self, p, indent, tone):
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(int(Pt(indent))))
        pPr.set("indent", str(-int(Pt(indent))))
        clr = etree.SubElement(pPr, qn("a:buClr"))
        etree.SubElement(clr, qn("a:srgbClr")).set("val", self.hex(tone).lstrip("#"))
        etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
        etree.SubElement(pPr, qn("a:buChar")).set("char", "•")

    # shapes --------------------------------------------------------------------
    @staticmethod
    def strip_style(shape):
        st = shape._element.find(qn("p:style"))
        if st is not None:
            shape._element.remove(st)

    def shape(self, slide, it):
        x, y, w, h = it["box"]
        kind = {"rect": MSO_SHAPE.RECTANGLE, "round": MSO_SHAPE.ROUNDED_RECTANGLE, "oval": MSO_SHAPE.OVAL}[it["t"]]
        sh = slide.shapes.add_shape(kind, Pt(x), Pt(y), Pt(w), Pt(h))
        self.strip_style(sh)
        if it.get("fill"):
            sh.fill.solid()
            sh.fill.fore_color.rgb = self.rgb(it["fill"])
        else:
            sh.fill.background()
        if it.get("line"):
            sh.line.color.rgb = self.rgb(it["line"])
            sh.line.width = Pt(it.get("lw", 0.75))
        else:
            sh.line.fill.background()
        if it["t"] == "round" and min(w, h) > 0:
            sh.adjustments[0] = min(0.5, it.get("radius", 8) / min(w, h))

    def line(self, slide, it):
        (x1, y1), (x2, y2) = it["p1"], it["p2"]
        c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
        self.strip_style(c)
        c.line.color.rgb = self.rgb(it.get("color", "line"))
        c.line.width = Pt(it.get("lw", 0.75))
        if it.get("dash"):
            c.line.dash_style = MSO_LINE.DASH
        if it.get("arrow"):
            ln = c.line._get_or_add_ln()
            tail = etree.SubElement(ln, qn("a:tailEnd"))
            tail.set("type", "triangle")
            tail.set("w", "med")
            tail.set("len", "med")

    # images --------------------------------------------------------------------
    def image(self, slide, it):
        from PIL import Image

        path: Path = it["path"]
        svg = None
        if path.suffix.lower() == ".svg":
            svg = self.themed_svg(path)
            if svg is None:
                return
            path = self.svg_to_png(svg)
            if path is None:
                self.ctx.err("visual", f"SVG를 PNG로 바꾸지 못했습니다: {svg.name}. PNG를 직접 넣으세요")
                return
        x, y, w, h = it["box"]
        with Image.open(path) as im:
            iw, ih = im.size
        if it.get("fit") == "cover":
            pic = slide.shapes.add_picture(str(path), Pt(x), Pt(y), Pt(w), Pt(h))
            box_r, img_r = w / h, iw / ih
            if img_r > box_r:
                cut = (1 - box_r / img_r) / 2
                pic.crop_left = pic.crop_right = cut
            else:
                cut = (1 - img_r / box_r) / 2
                pic.crop_top = pic.crop_bottom = cut
        else:
            scale = min(w / iw, h / ih)
            dw, dh = iw * scale, ih * scale
            pic = slide.shapes.add_picture(str(path), Pt(x + (w - dw) / 2), Pt(y + (h - dh) / 2), Pt(dw), Pt(dh))
        if svg is not None:
            self.attach_svg(slide, pic, svg.read_bytes())

    def themed_svg(self, svg: Path) -> Path | None:
        """Fill {{token}} colors in an SVG from the theme, so one drawing works in every theme."""
        import re

        text = svg.read_text(encoding="utf-8")
        tokens = set(re.findall(r"\{\{\s*([a-z_0-9]+)\s*\}\}", text))
        if not tokens:
            return svg
        colors = self.theme["colors"]
        unknown = tokens - set(colors)
        if unknown:
            self.ctx.err("visual", f"{svg.name}: 테마에 없는 색 이름 {', '.join(sorted(unknown))}")
            return None
        out = self.cache / f"{svg.stem}.{self.theme.get('name', 'theme')}.svg"
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(re.sub(r"\{\{\s*([a-z_0-9]+)\s*\}\}", lambda m: colors[m.group(1)], text), encoding="utf-8")
        return out

    def attach_svg(self, slide, pic, svg_bytes: bytes):
        from pptx.opc.constants import RELATIONSHIP_TYPE as RT
        from pptx.opc.package import Part

        pkg = slide.part.package
        partname = pkg.next_partname("/ppt/media/image%d.svg")
        part = Part(partname, "image/svg+xml", pkg, svg_bytes)
        rid = slide.part.relate_to(part, RT.IMAGE)
        blip = pic._element.find(".//" + qn("a:blip"))
        ext_lst = blip.find(qn("a:extLst"))
        if ext_lst is None:
            ext_lst = etree.SubElement(blip, qn("a:extLst"))
        ext = etree.SubElement(ext_lst, qn("a:ext"))
        ext.set("uri", SVG_EXT_URI)
        svg_blip = etree.SubElement(ext, f"{{{SVG_NS}}}svgBlip", nsmap={"asvg": SVG_NS})
        svg_blip.set(qn("r:embed"), rid)

    def logo(self, slide, it):
        logo = self.theme.get("logo")
        if not logo:
            return
        x, y, _, h = it["box"]
        name = logo.get(it.get("variant", "full"))
        for d in self.logo_dirs:
            p = d / name if name else None
            if p and p.is_file():
                slide.shapes.add_picture(str(p), Pt(x), Pt(y), height=Pt(h))
                return
        word = logo.get("wordmark")
        if not word:
            return
        # Fallback wordmark: the last letter set at cap height, then the signature dot.
        size = h * 0.95
        head, tail = word[:-1], word[-1]
        tb = slide.shapes.add_textbox(Pt(x), Pt(y - size * 0.25), Pt(size * 6), Pt(size * 1.4))
        tf = tb.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.word_wrap = False
        p = tf.paragraphs[0]
        for chunk, sz in ((head, size), (tail, size * 1.33)):
            r = p.add_run()
            r.text = chunk
            r.font.size = Pt(sz)
            r.font.bold = True
            r.font.color.rgb = self.rgb("text")
            self.set_font(r._r.get_or_add_rPr())
        m = self.ctx.measure
        wdt = m.width(head, size, True) + m.width(tail, size * 1.33, True)
        d = size * 0.2
        self.shape(slide, {"t": "oval", "box": (x + wdt + d * 0.6, y + size * 0.78, d, d), "fill": "accent"})

    # charts --------------------------------------------------------------------
    def number_format(self, v: dict) -> str:
        if v.get("format"):
            return v["format"]
        dec = int(v.get("decimals", 0))
        base = "#,##0" + ("." + "0" * dec if dec else "")
        unit = v.get("unit", "")
        return base + (f'"{unit}"' if unit else "")

    def chart_font(self, font, size=12, color="text2", bold=False):
        font.size = Pt(size)
        font.bold = bold
        font.color.rgb = self.rgb(color)
        self.set_font(font._rPr)

    @staticmethod
    def spPr(el):
        """c:spPr of a chartSpace (goes right after c:chart) or a plotArea (before c:extLst)."""
        sp = el.find(qn("c:spPr"))
        if sp is not None:
            return sp
        sp = etree.Element(qn("c:spPr"))
        anchor = el.find(qn("c:chart"))
        if anchor is not None:
            anchor.addnext(sp)
            return sp
        ext = el.find(qn("c:extLst"))
        if ext is not None:
            ext.addprevious(sp)
        else:
            el.append(sp)
        return sp

    def no_fill(self, el):
        sp = self.spPr(el)
        for child in list(sp):
            if child.tag in (qn("a:solidFill"), qn("a:noFill"), qn("a:gradFill")):
                sp.remove(child)
        sp.insert(0, etree.Element(qn("a:noFill")))
        ln = sp.find(qn("a:ln"))
        if ln is None:
            ln = etree.SubElement(sp, qn("a:ln"))
        for child in list(ln):
            ln.remove(child)
        etree.SubElement(ln, qn("a:noFill"))

    def chart(self, slide, it):
        v = it["spec"]
        x, y, w, h = it["box"]
        typ = v.get("type", "bar")
        kinds = {"bar": XL_CHART_TYPE.BAR_CLUSTERED, "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
                 "line": XL_CHART_TYPE.LINE_MARKERS}
        if typ not in kinds:
            self.ctx.err("visual", f"차트 type은 bar, column, line 중 하나입니다: {typ!r}")
            return
        cats = v["categories"]
        series = v["series"]
        fmt = self.number_format(v)
        data = CategoryChartData()
        data.categories = cats
        for sr in series:
            data.add_series(sr.get("name", ""), sr["values"], number_format=fmt)
        gf = slide.shapes.add_chart(kinds[typ], Pt(x), Pt(y), Pt(w), Pt(h), data)
        ch = gf.chart
        self.chart_font(ch.font, 12)
        ch.has_title = False
        self.no_fill(ch._chartSpace)
        self.no_fill(ch._chartSpace.chart.plotArea)
        multi = len(series) > 1
        ch.has_legend = multi
        if multi:
            ch.legend.position = XL_LEGEND_POSITION.TOP
            ch.legend.include_in_layout = False
            self.chart_font(ch.legend.font, 12)
        tone = v.get("tone", "accent")
        hl = v.get("highlight", [])
        hl = hl if isinstance(hl, list) else [hl]
        hl_idx = {cats.index(h_) if h_ in cats else h_ for h_ in hl if isinstance(h_, int) or h_ in cats}
        hl_series = {i for i, sr in enumerate(series) if sr.get("highlight")}
        values = [val for sr in series for val in sr["values"] if val is not None]
        vmax, vmin = max(values), min(values)
        va, ca = ch.value_axis, ch.category_axis
        va.has_minor_gridlines = False
        ca.has_major_gridlines = False
        ca.major_tick_mark = XL_TICK_MARK.NONE
        ca.format.line.color.rgb = self.rgb("line")
        self.chart_font(ca.tick_labels.font, 13 if typ == "bar" else 12, "text")
        plot = ch.plots[0]
        if typ in ("bar", "column"):
            va.visible = False
            va.has_major_gridlines = False
            va.minimum_scale = 0 if vmin >= 0 else vmin * 1.2
            va.maximum_scale = vmax * (1.22 if typ == "bar" else 1.15)
            if typ == "bar":
                ca.reverse_order = True
            plot.gap_width = 55
            plot.overlap = -10 if multi else 0
            for si, s in enumerate(plot.series):
                for pi in range(len(cats)):
                    on = (pi in hl_idx) if not multi else (si in hl_series)
                    col = tone if on else "chart_base"
                    pt = s.points[pi]
                    pt.format.fill.solid()
                    pt.format.fill.fore_color.rgb = self.rgb(col)
                    pt.format.line.fill.background()
                    if v.get("labels", True):
                        dl = pt.data_label
                        dl.number_format = fmt
                        dl.number_format_is_linked = False
                        dl.position = XL_LABEL_POSITION.OUTSIDE_END
                        self.chart_font(dl.font, 13, tone if on else "text2", bold=True)
        else:
            va.visible = True
            va.has_major_gridlines = True
            va.major_gridlines.format.line.color.rgb = self.rgb("line")
            va.format.line.fill.background()
            va.major_tick_mark = XL_TICK_MARK.NONE
            # A line shows change, so the axis starts near the data unless zero is asked for.
            lo, hi, step = nice_axis(0 if v.get("zero") else vmin, vmax)
            va.minimum_scale, va.maximum_scale, va.major_unit = lo, hi, step
            self.chart_font(va.tick_labels.font, 11, "muted")
            va.tick_labels.number_format = "#,##0"
            va.tick_labels.number_format_is_linked = False
            for si, s in enumerate(plot.series):
                on = si in hl_series or (not hl_series and si == 0)
                col = tone if on else "chart_base"
                s.smooth = False
                s.format.line.color.rgb = self.rgb(col)
                s.format.line.width = Pt(2.75 if on else 1.75)
                if not on:
                    s.format.line.dash_style = MSO_LINE.DASH
                s.marker.style = XL_MARKER_STYLE.CIRCLE
                s.marker.size = 6
                s.marker.format.fill.solid()
                s.marker.format.fill.fore_color.rgb = self.rgb(col)
                s.marker.format.line.color.rgb = self.rgb(col)
                last = series[si]["values"][-1]
                if last is not None and v.get("labels", True):
                    dl = s.points[len(cats) - 1].data_label
                    dl.has_text_frame = True
                    dl.text_frame.text = self.format_value(last, v)
                    dl.position = XL_LABEL_POSITION.RIGHT
                    for r in dl.text_frame.paragraphs[0].runs:
                        self.chart_font(r.font, 13, col if on else "text2", bold=True)

    @staticmethod
    def format_value(val, v):
        dec = int(v.get("decimals", 0))
        return f"{val:,.{dec}f}{v.get('unit', '')}"

    # tables --------------------------------------------------------------------
    def table(self, slide, it):
        v = it["spec"]
        x, y, w, _ = it["box"]
        header = v.get("header") or []
        rows = ([header] if header else []) + (v.get("rows") or [])
        ncol = len(it["col_w"])
        gf = slide.shapes.add_table(len(rows), ncol, Pt(x), Pt(y), Pt(w), Pt(sum(it["row_h"])))
        tbl = gf.table
        tbl_pr = gf._element.graphic.graphicData.tbl.tblPr
        sid = tbl_pr.find(qn("a:tableStyleId"))
        if sid is None:
            sid = etree.SubElement(tbl_pr, qn("a:tableStyleId"))
        sid.text = NO_TABLE_STYLE
        tbl.first_row = bool(header)
        tbl.horz_banding = False
        for ci, cw in enumerate(it["col_w"]):
            tbl.columns[ci].width = Pt(cw)
        for ri, rh in enumerate(it["row_h"]):
            tbl.rows[ri].height = Pt(rh)
        hl_row = v.get("highlight_row")
        hl_col = v.get("highlight_col")
        for ri, row in enumerate(rows):
            is_head = bool(header) and ri == 0
            data_ri = ri - (1 if header else 0)
            for ci in range(ncol):
                cell = tbl.cell(ri, ci)
                cell.margin_left = cell.margin_right = Pt(10)
                cell.margin_top = cell.margin_bottom = Pt(8)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf = cell.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                p.line_spacing = Pt(14 * 1.3)
                strong = is_head or data_ri == hl_row or ci == hl_col
                tone = "accent" if (not is_head and (data_ri == hl_row or ci == hl_col)) else "text"
                for chunk, bold, acc in parse_runs(str(row[ci] if ci < len(row) else "")):
                    r = p.add_run()
                    r.text = chunk
                    r.font.size = Pt(14)
                    r.font.bold = bool(strong or bold)
                    r.font.color.rgb = self.rgb("accent" if acc else tone)
                    self.set_font(r._r.get_or_add_rPr())
                if is_head:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = self.rgb("surface")
                else:
                    cell.fill.background()
                self.cell_borders(cell, bottom=("text" if is_head else "line"), width=1.0 if is_head else 0.75)

    def cell_borders(self, cell, bottom: str, width: float):
        tc_pr = cell._tc.get_or_add_tcPr()
        for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
            old = tc_pr.find(qn(tag))
            if old is not None:
                tc_pr.remove(old)
        for i, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
            ln = etree.Element(qn(tag))
            if tag == "a:lnB":
                ln.set("w", str(int(Pt(width))))
                fill = etree.SubElement(ln, qn("a:solidFill"))
                etree.SubElement(fill, qn("a:srgbClr")).set("val", self.hex(bottom).lstrip("#"))
            else:
                ln.set("w", "0")
                etree.SubElement(ln, qn("a:noFill"))
            tc_pr.insert(i, ln)
